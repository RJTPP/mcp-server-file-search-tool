from pathlib import Path
import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from typing import Union, List
import pymupdf4llm
import re

def remove_base64_images(markdown_text: str) -> str:
    # Match ![](data:image/<type>;base64,...)
    pattern = r'!\[\]\(data:image\/[a-zA-Z]+;base64,[^\)]*\)'
    cleaned_text = re.sub(pattern, '![Image]', markdown_text)
    return cleaned_text

def read_pdf(file_path: str, split_lines: bool = False, image_placeholder: bool = False) -> Union[str, List[str], List[List[str]]]:
    try:
        text = pymupdf4llm.to_markdown(
            file_path,
            embed_images=image_placeholder,
            image_format="jpg",
            image_path="assets/"
        )
        if image_placeholder:
            text = remove_base64_images(text)
        if split_lines:
            return [t for t in text.split("\n") if t.strip()]
        return text
    except Exception:
        doc = fitz.open(file_path)
        text = ""
        for page in doc:
            text += page.get_text()
    if split_lines:
        return [t for t in text.split("\n") if t.strip()]
    return text
    

def docx_iter_block_items(parent):
    """Yield paragraphs and tables in the order they appear."""
    for child in parent.element.body:
        if child.tag.endswith('}p'):
            yield Paragraph(child, parent)
        elif child.tag.endswith('}tbl'):
            yield Table(child, parent)


def read_docx(file_path: str, beautiful_table: bool = False, split_lines: bool = False) -> Union[str, List[str]]:
    doc = Document(file_path)
    text = ""
    table_placeholders: List[Tuple[str, Table]] = []
    table_idx = 0

    for block in docx_iter_block_items(doc):
        if isinstance(block, Paragraph):
            text += block.text + "\n"
        elif isinstance(block, Table):
            placeholder = f"__TABLE_PLACEHOLDER_{table_idx}__"
            table_placeholders.append((placeholder, block))
            text += f"{placeholder}\n"
            table_idx += 1

    if split_lines:
        result_lines = []
        for line in text.splitlines():
            found = False
            for placeholder, table in table_placeholders:
                if line.strip() == placeholder:
                    found = True
                    table_rows = ["<table>"]
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        row_text = " | ".join(cells) if beautiful_table else str(cells)
                        table_rows.append(row_text)
                    table_rows.append("</table>")
                    result_lines.append("\n".join(table_rows))  # ← One chunk per table
                    break
            if not found and line.strip():
                result_lines.append(line.strip())
        return result_lines
    else:
        for placeholder, table in table_placeholders:
            table_rows = ["<table>"]
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(cells) if beautiful_table else str(cells)
                table_rows.append(row_text)
            table_rows.append("</table>")
            formatted_table = "\n".join(table_rows)
            text = text.replace(placeholder, formatted_table)
        return text.strip()

    return text